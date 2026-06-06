import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO
import re


# ─── FILE READERS ─────────────────────────────────────────────────────────────

def read_dfa_report(file):
    df_raw = pd.read_excel(file, header=None, engine='openpyxl')

    hdr_idx = None
    for i, row in df_raw.iterrows():
        vals = [str(v).lower().strip() for v in row.values]
        if 'placement id' in vals and 'impressions' in vals:
            hdr_idx = i
            break
    if hdr_idx is None:
        raise ValueError("Could not find column headers in DFA report. Make sure you uploaded the correct file.")

    metadata = df_raw.iloc[:hdr_idx].copy()
    headers  = df_raw.iloc[hdr_idx].tolist()
    data     = df_raw.iloc[hdr_idx + 1:].copy()
    data.columns = headers
    data = data.reset_index(drop=True)

    pid_col = next((c for c in data.columns if str(c).lower().strip() == 'placement id'), None)
    if pid_col:
        data = data[pd.to_numeric(data[pid_col], errors='coerce').notna()].reset_index(drop=True)

    rename = {}
    for col in data.columns:
        cl = str(col).lower().strip()
        if cl == 'placement id':      rename[col] = 'placement_id'
        elif cl == 'placement':       rename[col] = 'placement_name'
        elif 'effective cpm' in cl:   rename[col] = 'cpm'
        elif cl == 'impressions':     rename[col] = 'impressions'
        elif 'media cost' in cl:      rename[col] = 'media_cost'
    data = data.rename(columns=rename)

    data['placement_id'] = data['placement_id'].apply(
        lambda x: str(int(float(x))) if pd.notna(x) else x
    )
    return metadata, data


def _is_line_col(val):
    """Match 'Line Item#', 'Line Item', 'Ad Book Line', 'AdBook Line', etc."""
    v = str(val).lower().strip()
    return bool(re.search(r'line.*(item|#)|ad\s*book\s*line', v))


def read_tag_sheet(file):
    xl = pd.ExcelFile(file, engine='openpyxl')

    # Collect placement→line mappings from ALL non-legacy sheets and merge.
    # Tag files often split placements across 'Tracking Ads' and 'Tags' sheets —
    # stopping at the first sheet would miss display placements in later sheets.
    all_dfs = []
    for sheet in xl.sheet_names:
        if 'legacy' in sheet.lower():
            continue
        df_raw = pd.read_excel(file, sheet_name=sheet, header=None, engine='openpyxl', nrows=30)
        chosen_hdr = None
        for i, row in df_raw.iterrows():
            vals = [str(v).lower().strip() for v in row.values]
            if 'placement id' in vals and any(_is_line_col(v) for v in vals):
                chosen_hdr = i
                break
        if chosen_hdr is None:
            continue

        df_raw  = pd.read_excel(file, sheet_name=sheet, header=None, engine='openpyxl')
        headers = df_raw.iloc[chosen_hdr].tolist()
        data    = df_raw.iloc[chosen_hdr + 1:].copy()
        data.columns = headers
        data = data.reset_index(drop=True)

        pid_col  = next((c for c in data.columns if str(c).lower().strip() == 'placement id'), None)
        line_col = next((c for c in data.columns if _is_line_col(c)), None)
        name_col = next((c for c in data.columns if 'placement name' in str(c).lower()), None)
        if not pid_col or not line_col:
            continue

        keep = [pid_col, line_col] + ([name_col] if name_col else [])
        df   = data[keep].dropna(subset=[pid_col]).copy()
        df.columns = ['placement_id', 'line_item'] + (['placement_name'] if name_col else [])

        df['placement_id'] = df['placement_id'].apply(
            lambda x: str(int(float(x))) if pd.notna(x) and str(x).replace('.', '').isdigit() else str(x).strip()
        )
        df['line_item'] = df['line_item'].apply(
            lambda x: str(int(float(x))) if pd.notna(x) and str(x).replace('.', '').isdigit() else str(x).strip()
        )
        all_dfs.append(df)

    if not all_dfs:
        raise ValueError(
            "Could not find a sheet with both 'Placement ID' and a Line Item column. "
            "Check that the correct tag mapping file was uploaded."
        )

    combined = pd.concat(all_dfs, ignore_index=True)
    combined = combined.drop_duplicates(subset=['placement_id']).reset_index(drop=True)
    return combined


def read_internal_billing(file, io_number):
    xl = pd.ExcelFile(file, engine='openpyxl')
    sheet = next((s for s in xl.sheet_names if '3pas' in s.lower()), None)
    if not sheet:
        raise ValueError("Cannot find a '3PAS Campaigns' sheet in the internal billing report.")

    df = pd.read_excel(file, sheet_name=sheet, header=1, engine='openpyxl')

    def clean_id(x):
        s = str(x).strip()
        return str(int(float(s))) if s.replace('.', '').isdigit() else s

    campaign_id_col = df.columns[3]
    filtered = df[df[campaign_id_col].apply(clean_id) == str(io_number).strip()].copy()

    if filtered.empty:
        raise ValueError(
            f"No rows found for IO number '{io_number}'. "
            "Confirm the Campaign ID matches Column D of the billing report."
        )

    result = pd.DataFrame({
        'month':               filtered.iloc[:, 0].astype(str).str.strip(),   # Col A
        'line_item':           filtered.iloc[:, 6].apply(clean_id),
        'campaign_name':       filtered.iloc[:, 4].astype(str).str.strip(),
        'position_path':       filtered.iloc[:, 5].astype(str).str.strip(),
        'impressions_sold':    pd.to_numeric(filtered.iloc[:, 20], errors='coerce').fillna(0),
        'discrepancy_flag':    filtered.iloc[:, 34].astype(str).str.strip(),
        'final_billable_impr': pd.to_numeric(filtered.iloc[:, 35], errors='coerce').fillna(0),
        'final_invoicing_amt': pd.to_numeric(filtered.iloc[:, 36], errors='coerce').fillna(0),
    })
    return result


def build_filename(campaign_name, month):
    """e.g. 'Visit Loudoun VA Spring 2026 Campaign', 'April' → 'VisitLoudounVASpring2026-April Billables.xlsx'"""
    clean = re.sub(r'\bCampaign\b', '', campaign_name, flags=re.IGNORECASE)
    clean = re.sub(r'[^A-Za-z0-9]', '', clean)   # remove spaces and special chars
    return f"{clean}-{month} Billables.xlsx"


# ─── BILLING LOGIC ────────────────────────────────────────────────────────────

def is_first_party(position_path):
    p = str(position_path).lower()
    return 'mobile app' in p or 'interstitial' in p


def process_campaign(dfa_meta, dfa_df, tag_df, billing_df, notes):
    warnings = []

    # DFA lookup
    dfa_lookup = {}
    for _, row in dfa_df.iterrows():
        pid = str(row['placement_id']).strip()
        dfa_lookup[pid] = {
            'name':        str(row.get('placement_name', '')).strip(),
            'impressions': int(row.get('impressions', 0) or 0),
            'cpm':         float(row.get('cpm', 0) or 0),
            'media_cost':  float(row.get('media_cost', 0) or 0),
        }

    # Tag lookup: line → [pids]
    line_to_pids = {}
    for _, row in tag_df.iterrows():
        line_to_pids.setdefault(str(row['line_item']), []).append(str(row['placement_id']))

    # DFA placements not in tag sheet (e.g. 1x1 impression trackers)
    tagged_pids  = {str(r['placement_id']) for _, r in tag_df.iterrows()}
    untagged_dfa = {pid: info for pid, info in dfa_lookup.items() if pid not in tagged_pids}

    # For DFA placements missing from the tag sheet, infer line from placement name
    # (e.g. "...Line Item 2..." → line 2). Handles placements added after tag sheet was generated.
    for pid, info in untagged_dfa.items():
        m = re.search(r'line\s*item\s*(\d+)', info['name'], re.IGNORECASE)
        if m:
            line_to_pids.setdefault(str(int(m.group(1))), []).append(pid)

    rows_3p = []
    rows_1p = []
    verification = []

    billing_sorted = billing_df.copy()
    billing_sorted['_line_int'] = pd.to_numeric(billing_sorted['line_item'], errors='coerce').fillna(999)
    billing_sorted = billing_sorted.sort_values('_line_int').reset_index(drop=True)

    for _, bill in billing_sorted.iterrows():
        line     = str(bill['line_item']).strip()
        flag     = str(bill['discrepancy_flag']).strip()
        capped   = int(bill['impressions_sold'])
        aj       = float(bill['final_billable_impr'])
        ak       = float(bill['final_invoicing_amt'])
        pos_path = str(bill['position_path']).strip()
        line_int = int(line) if line.isdigit() else line

        # Hard stop: variance flag
        if '10%' in flag or ('variance' in flag.lower() and 'above' in flag.lower()):
            warnings.append(
                f"🚨 LINE {line_int}: Discrepancy flagged — \"{flag}\". "
                "Resolve before billing this line."
            )
            continue

        is_free = 'free line' in flag.lower()

        # 1st party: Mobile App or Interstitial
        if is_first_party(pos_path):
            cpm_1p = round(ak / (aj / 1000), 2) if aj > 0 else 0
            rows_1p.append({
                'placement_id':   None,
                'placement_name': pos_path,
                'line':           line_int,
                'impressions':    int(aj),
                'capped':         capped,
                'cpm':            cpm_1p,
                'media_cost':     ak,
                '_first_in_line': True,
            })
            verification.append({
                'Line':            line_int,
                'Position':        pos_path,
                'DFA Impressions': 'N/A — 1st Party',
                'Internal AJ':     f"{int(aj):,}",
                'Difference':      '—',
                'Billed $':        f"${ak:,.3f}",
                'Status':          '✅ OK — 1st Party (Mobile App / Interstitial)',
            })
            continue

        # 3rd party: match placements via tag sheet
        line_pids = line_to_pids.get(line, [])
        matched   = [pid for pid in line_pids if pid in dfa_lookup]

        # Fallback: match untagged DFA placements by name keyword similarity
        if not matched:
            tag_names = []
            if 'placement_name' in tag_df.columns:
                tag_names = tag_df[tag_df['line_item'] == line]['placement_name'].dropna().tolist()
            keywords = set()
            for tn in tag_names:
                for part in re.split(r'[_|/]', str(tn).lower()):
                    if len(part.strip()) > 4:
                        keywords.add(part.strip())
            for pid, info in untagged_dfa.items():
                nm = info['name'].lower()
                if sum(1 for k in keywords if k in nm) >= 3:
                    matched.append(pid)

        # Still no match: fall back to AJ/AK, treat as 1P
        if not matched and not is_free:
            warnings.append(
                f"ℹ️ LINE {line_int}: No DFA placements matched. "
                "Falling back to internal AJ/AK."
            )
            cpm_fb = round(ak / (aj / 1000), 2) if aj > 0 else 0
            rows_1p.append({
                'placement_id':   None,
                'placement_name': pos_path,
                'line':           line_int,
                'impressions':    int(aj),
                'capped':         capped,
                'cpm':            cpm_fb,
                'media_cost':     ak,
                '_first_in_line': True,
            })
            verification.append({
                'Line':            line_int,
                'Position':        pos_path,
                'DFA Impressions': 'No DFA match found',
                'Internal AJ':     f"{int(aj):,}",
                'Difference':      '—',
                'Billed $':        f"${ak:,.3f}",
                'Status':          '⚠️ Fallback — No DFA match (review manually)',
            })
            continue

        # Compute line totals
        dfa_total_impr = sum(dfa_lookup[p]['impressions'] for p in matched)
        dfa_total_cost = 0.0 if is_free else sum(dfa_lookup[p]['media_cost'] for p in matched)

        # Variance check
        if is_free:
            status   = '✅ Free Line-OK — $0'
            diff_str = '—'
        else:
            diff     = dfa_total_impr - int(aj)
            pct      = abs(diff) / max(int(aj), 1) * 100
            diff_str = f"{diff:+,}"
            if pct > 10:
                warnings.append(
                    f"🚨 LINE {line_int}: DFA ({dfa_total_impr:,}) vs Internal AJ ({int(aj):,}) "
                    f"= {pct:.1f}% variance. Investigate before billing."
                )
                status = f'🚨 VARIANCE {pct:.1f}% — DO NOT BILL'
            else:
                status = f'✅ OK — Using DFA  ({diff:+,} vs AJ)'

        verification.append({
            'Line':            line_int,
            'Position':        pos_path,
            'DFA Impressions': f"{dfa_total_impr:,}" + (' (Free)' if is_free else ''),
            'Internal AJ':     f"{int(aj):,}",
            'Difference':      diff_str,
            'Billed $':        f"${dfa_total_cost:,.3f}",
            'Status':          status,
        })

        first = True
        for pid in matched:
            info = dfa_lookup[pid]
            rows_3p.append({
                'placement_id':   int(pid) if pid.isdigit() else pid,
                'placement_name': info['name'],
                'line':           line_int if first else None,
                'impressions':    info['impressions'],
                'capped':         capped if first else None,
                'cpm':            info['cpm'],
                'media_cost':     0.0 if is_free else info['media_cost'],
                '_first_in_line': first,
            })
            first = False

    # Cross-check totals: DFA raw total + 1P AK total should equal output grand total
    dfa_raw_total      = sum(info['media_cost'] for info in dfa_lookup.values())
    internal_1p_total  = billing_df[billing_df['position_path'].apply(is_first_party)]['final_invoicing_amt'].sum()

    # 3P lines first (sorted), then 1P lines at end
    return rows_3p + rows_1p, verification, warnings, dfa_raw_total, internal_1p_total


# ─── EXCEL GENERATOR ──────────────────────────────────────────────────────────

LINE_COLORS = [
    'DEEAF1', 'E2EFDA', 'FFF2CC', 'FCE4D6',
    'EAD1DC', 'F4CCCC', 'D9EAD3', 'CFE2F3',
]


def generate_excel(output_rows, dfa_meta):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Billables"

    # DFA metadata rows
    for r_idx, row in dfa_meta.iterrows():
        for c_idx, val in enumerate(row, 1):
            if pd.notna(val) and str(val).strip() not in ('', 'nan'):
                cell = ws.cell(row=r_idx + 1, column=c_idx, value=val)
                cell.font = Font(bold=(r_idx == 0), size=12 if r_idx == 0 else 9)

    meta_rows = len(dfa_meta)
    lbl_row   = meta_rows + 1
    hdr_row   = lbl_row + 1

    # Report Fields / Metrics label row
    ws.cell(row=lbl_row, column=1, value="Report Fields").font = Font(bold=True, size=9)
    ws.cell(row=lbl_row, column=4, value="Report Metrics").font = Font(bold=True, size=9)

    # Column headers
    HDR  = ["Placement ID", "Placement", "Line", "Impressions", "Capped Impressions total", "CPM", "Media Cost"]
    hfil = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    hfnt = Font(bold=True, color="FFFFFF", size=10)
    for ci, h in enumerate(HDR, 1):
        c = ws.cell(row=hdr_row, column=ci, value=h)
        c.fill = hfil; c.font = hfnt
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[hdr_row].height = 30

    # Data rows
    line_color_map = {}
    color_idx      = 0
    current_line   = None

    for r_off, row in enumerate(output_rows):
        r = hdr_row + 1 + r_off
        if row['line'] is not None:
            current_line = row['line']
        if current_line not in line_color_map:
            line_color_map[current_line] = LINE_COLORS[color_idx % len(LINE_COLORS)]
            color_idx += 1
        fill = PatternFill(
            start_color=line_color_map[current_line],
            end_color=line_color_map[current_line],
            fill_type="solid"
        )
        vals = [
            row['placement_id'], row['placement_name'], row['line'],
            row['impressions'],  row['capped'],         row['cpm'],  row['media_cost'],
        ]
        for ci, val in enumerate(vals, 1):
            cell = ws.cell(row=r, column=ci, value=val)
            cell.fill = fill
            cell.font = Font(size=9)
            cell.alignment = Alignment(vertical="center", wrap_text=(ci == 2))
            if ci == 1 and val is not None: cell.number_format = '0'
            if ci in (4, 5) and val is not None: cell.number_format = '#,##0'
            if ci == 6 and val is not None:       cell.number_format = '$#,##0.00'
            if ci == 7 and val is not None:       cell.number_format = '$#,##0.000'

    # Grand Total row
    total_row  = hdr_row + 1 + len(output_rows)
    total_cost = sum(r.get('media_cost', 0) or 0 for r in output_rows)
    tfil = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    tfnt = Font(bold=True, color="FFFFFF", size=10)
    for ci in range(1, 8):
        cell = ws.cell(row=total_row, column=ci)
        cell.fill = tfil; cell.font = tfnt
    ws.cell(row=total_row, column=1, value="Grand Total:").alignment = Alignment(horizontal="right")
    c7 = ws.cell(row=total_row, column=7, value=total_cost)
    c7.number_format = '$#,##0.000'; c7.fill = tfil; c7.font = tfnt

    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 85
    ws.column_dimensions['C'].width = 7
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 24
    ws.column_dimensions['F'].width = 10
    ws.column_dimensions['G'].width = 14

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ─── STREAMLIT UI ─────────────────────────────────────────────────────────────

st.set_page_config(
    page_title="3PAS Billing Tool",
    layout="wide",
    page_icon="📊",
)

st.title("📊 3PAS Monthly Billing Report Generator")
st.caption(
    "Generate client-facing billables reports from your DFA report, "
    "tag mapping, and internal billing files."
)
st.markdown("---")

# Step 1 ── Internal Billing Report
st.subheader("Step 1 — Internal Billing Report")
st.caption(
    "Upload the monthly consolidated billing report. "
    "This single file covers all campaigns you process below."
)
billing_file = st.file_uploader(
    "Internal Billing Report (.xlsx)",
    type=["xlsx"],
    key="billing",
    label_visibility="collapsed",
)

st.markdown("---")

# Step 2 ── Campaigns
st.subheader("Step 2 — Campaigns")
st.caption("Add one section per campaign. All campaigns run in a single click.")

if "num_campaigns" not in st.session_state:
    st.session_state.num_campaigns = 1

campaigns = []
for i in range(st.session_state.num_campaigns):
    with st.expander(f"**Campaign {i + 1}**", expanded=True):
        io_num = st.text_input(
            "IO Number (Campaign ID)",
            key=f"io_{i}",
            placeholder="e.g. 637466  —  used to filter the billing report",
        )
        col1, col2 = st.columns(2)
        with col1:
            st.markdown("**DFA Report**")
            dfa_file = st.file_uploader(
                "DFA Report", type=["xlsx"], key=f"dfa_{i}", label_visibility="collapsed"
            )
        with col2:
            st.markdown("**Tag Mapping**")
            tag_file = st.file_uploader(
                "Tag Mapping", type=["xlsx"], key=f"tag_{i}", label_visibility="collapsed"
            )
        notes = st.text_area(
            "Notes / Nuances",
            key=f"notes_{i}",
            placeholder="Any special instructions or flags for this campaign...",
            height=75,
        )
        campaigns.append({"io": io_num, "dfa": dfa_file, "tag": tag_file, "notes": notes, "idx": i + 1})

c_add, c_rem, _ = st.columns([1, 1, 6])
with c_add:
    if st.button("＋ Add Campaign"):
        st.session_state.num_campaigns += 1
        st.rerun()
if st.session_state.num_campaigns > 1:
    with c_rem:
        if st.button("－ Remove Last"):
            st.session_state.num_campaigns -= 1
            st.rerun()

st.markdown("---")

# Step 3 ── Run
run = st.button("▶  Run Billing", type="primary", use_container_width=True)

if run:
    errors = []
    if not billing_file:
        errors.append("Upload the Internal Billing Report in Step 1.")
    for c in campaigns:
        if not c["io"].strip():
            errors.append(f"Campaign {c['idx']}: IO Number is required.")
        if not c["dfa"]:
            errors.append(f"Campaign {c['idx']}: DFA Report is required.")
        if not c["tag"]:
            errors.append(f"Campaign {c['idx']}: Tag Mapping is required.")

    if errors:
        for e in errors:
            st.error(e)
    else:
        st.markdown("---")
        st.header("Results")

        for c in campaigns:
            st.markdown(f"### Campaign {c['idx']} — IO {c['io']}")
            try:
                with st.spinner(f"Processing Campaign {c['idx']}..."):
                    dfa_meta, dfa_data = read_dfa_report(c["dfa"])
                    tag_data           = read_tag_sheet(c["tag"])
                    billing_file.seek(0)
                    billing_data       = read_internal_billing(billing_file, c["io"])

                    output_rows, verification, warnings, dfa_raw_total, internal_1p_total = process_campaign(
                        dfa_meta, dfa_data, tag_data, billing_data, c["notes"]
                    )

                camp_name = billing_data["campaign_name"].iloc[0] if not billing_data.empty else ""
                if camp_name:
                    st.caption(f"**{camp_name}**")

                if c["notes"].strip():
                    st.info(f"📝 Notes: {c['notes'].strip()}")

                has_blocker = False
                for w in warnings:
                    if "🚨" in w:
                        st.error(w)
                        has_blocker = True
                    else:
                        st.warning(w)

                st.markdown("**Billing Verification**")
                st.dataframe(pd.DataFrame(verification), use_container_width=True, hide_index=True)

                total      = sum(r.get("media_cost", 0) or 0 for r in output_rows)
                cross_check = dfa_raw_total + internal_1p_total
                match       = abs(cross_check - total) < 0.01

                st.metric("Total Billed", f"${total:,.3f}")

                with st.expander("🔢 Cross-Check Verification", expanded=True):
                    col1, col2, col3 = st.columns(3)
                    col1.metric("DFA Report Total (3P)", f"${dfa_raw_total:,.3f}")
                    col2.metric("1P Internal AK Total", f"${internal_1p_total:,.3f}")
                    col3.metric("Cross-Check Sum", f"${cross_check:,.3f}")
                    if match:
                        st.success("✅ Cross-check passed — DFA total + 1P AK total matches output grand total.")
                    else:
                        st.error(f"⚠️ Cross-check mismatch — expected ${cross_check:,.3f}, got ${total:,.3f}. A 3P line may have fallen back to 1P or been misrouted.")

                if has_blocker:
                    st.error("⛔ Download blocked — resolve the flagged variance(s) above before billing.")
                else:
                    excel_bytes = generate_excel(output_rows, dfa_meta)
                    month       = billing_data["month"].iloc[0] if not billing_data.empty else "Unknown"
                    camp_nm     = billing_data["campaign_name"].iloc[0] if not billing_data.empty else c["io"]
                    filename    = build_filename(camp_nm, month)
                    st.download_button(
                        label="⬇  Download Billables Report",
                        data=excel_bytes,
                        file_name=filename,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary",
                    )

            except Exception as e:
                st.error(f"Error processing Campaign {c['idx']}: {str(e)}")

            st.markdown("---")
